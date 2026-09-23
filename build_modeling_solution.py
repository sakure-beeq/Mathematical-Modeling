from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path("/Users/wangyiming/Desktop/E题/复杂场景下多模态情感预测建模详细解题方案.docx")

BODY_CN = "Arial Unicode MS"
HEAD_CN = "Arial Unicode MS"
LATIN = "Arial Unicode MS"
NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F7"
GRID = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_run_font(run, name=BODY_CN, size=11, bold=None, color=BLACK, italic=None):
    run.font.name = LATIN
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), LATIN)
    rfonts.set(qn("w:hAnsi"), LATIN)
    rfonts.set(qn("w:eastAsia"), name)


def set_style_font(style, name, size, bold=False):
    style.font.name = LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), LATIN)
    rfonts.set(qn("w:hAnsi"), LATIN)
    rfonts.set(qn("w:eastAsia"), name)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def add_body(doc, text, *, bold_lead=None, indent=True, keep=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.32
    if indent:
        pf.first_line_indent = Cm(0.74)
    if keep:
        pf.keep_with_next = True
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.22
        for run in p.runs:
            set_run_font(run)
        if not p.runs:
            set_run_font(p.add_run(item))
        else:
            p.runs[0].text = item
            set_run_font(p.runs[0])


def add_numbered(doc, items):
    for idx, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        r = p.add_run(f"{idx}.  {item}")
        set_run_font(r)


def add_heading(doc, text, level=1, page_break=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    # 封面后已显式分页，其余章节自然衔接，避免产生大面积空白页。
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_run_font(r, HEAD_CN, {1: 16, 2: 13.5, 3: 12}[level], bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=10.5, bold=True)
    return p


def add_table(doc, headers, rows, widths=None, aligns=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 110, 120, 110, 120)
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(h))
        set_run_font(r, HEAD_CN, font_size, bold=True, color=WHITE)
    for ridx, row in enumerate(rows):
        tr = table.add_row()
        prevent_row_split(tr)
        for i, value in enumerate(row):
            cell = tr.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 95, 110, 95, 110)
            if widths:
                set_cell_width(cell, widths[i])
            if ridx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            align = aligns[i] if aligns else (WD_ALIGN_PARAGRAPH.LEFT if len(str(value)) > 14 else WD_ALIGN_PARAGRAPH.CENTER)
            p.alignment = align
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    return table


def add_omml_equation(doc, text, number=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    mrpr = OxmlElement("m:rPr")
    sty = OxmlElement("m:sty")
    sty.set(qn("m:val"), "p")
    mrpr.append(sty)
    mt = OxmlElement("m:t")
    mt.text = text
    mr.append(mrpr)
    mr.append(mt)
    omath.append(mr)
    omath_para.append(omath)
    p._p.append(omath_para)
    if number:
        r = p.add_run(f"    {number}")
        set_run_font(r, size=10.5)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9)


def configure_styles(doc):
    styles = doc.styles
    set_style_font(styles["Normal"], BODY_CN, 11, False)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.32
    set_style_font(styles["Title"], HEAD_CN, 22, True)
    styles["Title"].paragraph_format.space_after = Pt(15)
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    set_style_font(styles["Subtitle"], BODY_CN, 12, False)
    for i, size in ((1, 16), (2, 13.5), (3, 12)):
        set_style_font(styles[f"Heading {i}"], HEAD_CN, size, True)
        styles[f"Heading {i}"].paragraph_format.space_before = Pt(13 if i == 1 else 9)
        styles[f"Heading {i}"].paragraph_format.space_after = Pt(6)
        styles[f"Heading {i}"].paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2", "List Number"):
        set_style_font(styles[name], BODY_CN, 11, False)


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)
    add_page_number(section.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(105)
    r = p.add_run("复杂场景下多模态情感预测建模详细解题方案")
    set_run_font(r, HEAD_CN, 22, bold=True)
    p2 = doc.add_paragraph(style="Subtitle")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("数据综述 预处理方法 数学模型 实验设计与提交方案")
    set_run_font(r, BODY_CN, 12)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(35)
    r = p3.add_run("适用于 2026 年中国研究生数学建模竞赛 E 题")
    set_run_font(r, size=11)
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(160)
    r = p4.add_run("建模方案文档")
    set_run_font(r, HEAD_CN, 12, bold=True)
    p4.add_run().add_break(WD_BREAK.PAGE)

    add_heading(doc, "方案总览", 1)
    add_body(doc, "本题应沿用一条统一技术路线。问题一建立可追溯的三模态特征与共同时间轴；问题二在共同特征上加入连续区间缺失模拟、跨模态表示重建和动态可靠性融合；问题三复用缺失感知模型，并用模态 Shapley 值和时间片遮挡定位关键证据。主模型建议使用 aligned_50.pkl，对齐版结构简单、接口稳定，便于实现缺失掩码和解释定位。unaligned_50.pkl 适合作为扩展对比，不宜直接作为唯一主线。")
    add_caption(doc, "表 1 三个问题的建模主线")
    add_table(
        doc,
        ["问题", "核心任务", "推荐方法", "主要输出"],
        [
            ["问题一", "特征提取和时间对齐", "BERT 加声学特征加面部行为特征，采用词级区间对齐", "100 条样本特征、时间戳、掩码和日志"],
            ["问题二", "局部模态缺失下预测", "缺失感知多模态 Transformer、连续块遮挡、潜在重建、动态门控", "极性、强度、鲁棒性曲线和附件三预测"],
            ["问题三", "可量化和可复核解释", "精确模态 Shapley 值、连续窗口遮挡、证据删除验证", "模态贡献、关键文本、语音时段、视觉关键帧"],
        ],
        widths=[1.5, 3.2, 5.2, 3.7],
        font_size=9.2,
    )
    add_body(doc, "关键结论是先解决数据掩码，再讨论模型复杂度。附件二本身含有视觉全零和内部无效帧，附件三又引入连续缺失。如果把填充、质量失效和人工缺失混为一个零值条件，模型训练和缺失规律分析都会产生偏差。", bold_lead="关键结论是")

    add_heading(doc, "内容结构", 2)
    add_numbered(doc, [
        "题目目标与建模关系",
        "数据综述与质量核查",
        "统一数据预处理",
        "问题一特征提取与时序对齐",
        "问题二缺失条件下鲁棒预测",
        "问题三可解释性预测",
        "实验设计与评价方法",
        "结果文件与论文组织",
        "实施顺序与风险控制",
    ])

    add_heading(doc, "一 题目目标与建模关系", 1, page_break=True)
    add_body(doc, "题目要求依次完成特征提取与对齐、局部模态缺失下预测、完整模态条件下解释。三个问题并非独立任务。问题一决定统一输入结构和时间映射，问题二提供能处理模态子集的预测模型，问题三正好可以利用该模型计算不同模态组合的边际贡献。因此，最稳妥的方案是让三个问题共享数据接口、模态掩码和预测头。")
    add_caption(doc, "表 2 统一建模流程")
    add_table(
        doc,
        ["阶段", "输入", "处理", "输出"],
        [
            ["数据核查", "视频、标签和特征文件", "核对 ID、标签、维度、长度和零向量", "质量报告和异常清单"],
            ["特征构建", "文本、音频和视频", "提取特征并映射到共同时间区间", "固定长度特征与三类掩码"],
            ["鲁棒训练", "完整样本和模拟缺失样本", "重建缺失表示并动态融合", "分类头和回归头"],
            ["影响分析", "配对缺失实验", "比较缺失类型、比例和位置", "性能下降规律和置信区间"],
            ["解释生成", "完整样本和训练后模型", "模态 Shapley 值与窗口遮挡", "解释卡和关键证据定位"],
        ],
        widths=[2.1, 3.2, 5.0, 3.3],
    )
    add_body(doc, "评价包括分类 Accuracy 和 F1，以及回归 MAE 和 Pearson 相关系数。F1 的平均方式在题面中未明确，正文应将 Macro F1 作为主要结果，并同时列出 Weighted F1，避免类别分布不均衡掩盖中性类和负向类性能。")

    add_heading(doc, "二 数据综述与质量核查", 1, page_break=True)
    add_heading(doc, "附件一原始视频", 2)
    add_body(doc, "附件一实际包含 37 个 video_id 文件夹和 100 个视频片段。标签表共有 100 条记录，视频与标签通过 video_id 和 clip_id 完全一一对应，没有缺失文件或无标签文件。题面给出的时长范围为 2.648 秒至 34.567 秒。")
    add_caption(doc, "表 3 附件一标签分布")
    add_table(doc, ["类别", "样本数", "比例", "连续标签定义"], [
        ["Positive", "57", "57%", "大于 0"],
        ["Neutral", "25", "25%", "等于 0"],
        ["Negative", "18", "18%", "小于 0"],
    ], widths=[3.0, 2.5, 2.5, 5.6])
    add_body(doc, "连续标签的实际范围为 -2.000 至 2.667，均值约为 0.340，样本整体偏正向。英文文本长度为 5 至 65 个词，平均 19.32 个词。Excel 中的 label 列按字符串存储，读取后必须显式转换为浮点数。")

    add_heading(doc, "附件二标准化特征", 2)
    add_caption(doc, "表 4 附件二数据划分")
    add_table(doc, ["划分", "总数", "Positive", "Neutral", "Negative"], [
        ["train", "3395", "1670", "758", "967"],
        ["valid", "728", "338", "184", "206"],
        ["test", "727", "362", "158", "207"],
        ["合计", "4850", "2370", "1100", "1380"],
    ], widths=[2.5, 2.2, 2.8, 2.8, 2.8])
    add_body(doc, "总体类别比例约为正向 48.9%、负向 28.5%、中性 22.7%。连续标签范围为 -3 至 3，均值约为 0.161。train、valid 和 test 之间没有重复 video_id，因此没有同一原始视频跨集合泄漏的问题。")
    add_caption(doc, "表 5 两类特征文件的结构")
    add_table(doc, ["版本", "文本特征", "语音特征", "视觉特征", "时间结构"], [
        ["aligned_50", "50 × 768", "50 × 74", "50 × 35", "三模态共享 50 个位置"],
        ["unaligned_50", "50 × 768", "500 × 74", "500 × 35", "语音和视觉保留独立时间轴"],
    ], widths=[2.5, 2.5, 2.5, 2.5, 3.7])

    add_heading(doc, "数据质量发现", 2)
    add_caption(doc, "表 6 需要在建模前处理的数据问题")
    add_table(doc, ["发现", "核查结果", "建模影响", "处理原则"], [
        ["文本填充", "text 的填充向量不一定为零", "不能用 text 等于零判断有效长度", "以 text_bert 的 attention mask 为准"],
        ["文本截断", "352 条样本达到 50 个 token，148 条文本超过 48 个词", "尾部语义可能被截断", "记录截断标志并做敏感性分析"],
        ["全零视觉", "aligned 版 153 条视觉特征全部为零", "视觉模态在部分样本完全不可用", "设置质量掩码并允许退化为双模态"],
        ["内部视觉零行", "约 605 条在名义有效区间内出现零行", "人脸检测失败可能被误认为情绪信号", "区分质量失效与 padding"],
        ["未对齐长度异常", "648 条 vision_lengths 与最后非零位置不一致", "盲目信任长度字段会错误截断", "联合长度字段和非零位置生成掩码"],
    ], widths=[2.5, 3.2, 4.0, 3.9], font_size=9.0)
    add_body(doc, "上述结果支持将 aligned_50 作为主版本。该版本虽然也存在视觉无效位置，但长度固定、文本掩码明确，且附件三和附件四都提供同接口版本。未对齐版保留更细时间分辨率，适合在主模型完成后作为扩展对照。")

    add_heading(doc, "附件三缺失专项集", 2)
    add_body(doc, "附件三包含 30 个对齐文件和 30 个未对齐文件，两组表示同一批无标签样本。对齐版实际只包含 text_bert、audio 和 vision，样本 ID 应由文件名取得。text_bert 在该附件中为浮点型，输入 BERT 前必须转换为整数 token、attention mask 和 segment ID。30 条对齐样本中有 27 条在音频或视觉有效区间内出现内部全零片段，缺失主要集中在音频和视觉。")

    add_heading(doc, "附件四解释专项集", 2)
    add_body(doc, "附件四包含 20 条无标签样本，每条同时提供原始视频、对齐特征、未对齐特征和英文转写。数据核查发现，对齐版 13 号样本的视觉矩阵全部为零，未对齐版个别样本的 vision_lengths 也存在异常。论文应表述为原始三模态素材完整，但特征提取结果可能发生视觉质量失效，并由质量掩码处理。")

    add_heading(doc, "三 统一数据预处理", 1, page_break=True)
    add_body(doc, "数据预处理的核心是把序列填充、模态可用性和特征质量分开表示。设模态 m 在位置 t 的三类掩码分别为 P、A 和 Q，则模型实际使用的掩码为三者乘积。")
    add_omml_equation(doc, "Mₘ,ₜ = Pₘ,ₜ × Aₘ,ₜ × Qₘ,ₜ", "式 1")
    add_bullets(doc, [
        "P 表示真实序列位置，用于排除尾部填充。",
        "A 表示该位置的模态是否可用，用于描述题目设置的局部缺失。",
        "Q 表示特征质量是否合格，用于描述人脸检测失败、音频无效等问题。",
    ])
    add_caption(doc, "表 7 各类数据的预处理规则")
    add_table(doc, ["对象", "类型转换", "有效位置判断", "标准化和异常处理"], [
        ["分类标签", "转换为 int64", "0 负向 1 中性 2 正向", "仅使用 train 统计类别权重"],
        ["回归标签", "转换为 float32", "范围检查为 -3 至 3", "不缩放或线性缩放后反变换"],
        ["text_bert", "token 和 mask 转为 int64", "使用 attention mask", "统一冻结 BERT 检查点生成表示"],
        ["text", "转换为 float32", "不得用零值推断长度", "作为提供特征的对照版本"],
        ["audio", "float64 转 float32", "padding mask 加零行检查", "只在 train 有效位置计算均值和方差"],
        ["vision", "float64 转 float32", "padding mask 加质量 mask", "全零样本允许退化为双模态"],
    ], widths=[2.3, 3.0, 4.1, 4.2], font_size=9.1)
    add_body(doc, "标准化参数只能由训练集的有效位置计算。对任一模态，先在 M 等于 1 的位置估计均值和标准差，再进行 z-score 标准化。标准化完成后，将无效位置重新置零，防止原零向量减去均值后变成非零。可根据训练集分布对极端值进行正负 5 个标准差截尾。")
    add_heading(doc, "文本接口统一", 2)
    add_body(doc, "附件二含有 text 和 text_bert，但附件三对齐版没有 text。主模型应以 text_bert 为统一接口，使用同一个冻结 BERT 检查点重新生成文本表示。这样可以避免训练时使用预计算 text、推理时临时编码 text_bert 所造成的表示分布不一致。附件二提供的 text 可以用于核查或对照实验。")
    add_heading(doc, "数据泄漏控制", 2)
    add_bullets(doc, [
        "train 用于模型参数学习和标准化统计。",
        "valid 用于模型结构、超参数和分类阈值选择。",
        "test 仅用于最终性能报告，不参与阈值和结构选择。",
        "附件三和附件四没有标签，只进行最终推理，不能参与任何调参。",
    ])

    add_heading(doc, "四 问题一特征提取与时序对齐", 1, page_break=True)
    add_heading(doc, "总体方法", 2)
    add_body(doc, "问题一采用文本词级时间区间作为主时间轴。标签表已经提供英文转写，可对视频音轨和既有文本进行强制对齐，得到每个词的起止时间。文本、语音和视觉特征都在相同词区间内聚合，从而保证序列位置可以回溯到原始语句、音频时段和视频帧。")
    add_caption(doc, "表 8 三模态特征定义")
    add_table(doc, ["模态", "推荐工具或方法", "特征内容", "建议维度"], [
        ["文本", "英文 tokenizer 加冻结 BERT", "上下文词元表示，WordPiece 聚合为词级表示", "768"],
        ["语音", "COVAREP 或同类声学工具", "基频、浊音、能量、谱特征、MFCC 和韵律", "74"],
        ["视觉", "OpenFace 或同类面部分析工具", "动作单元、头部姿态、注视和面部关键点统计", "35"],
    ], widths=[2.0, 4.1, 5.3, 2.2], font_size=9.2)
    add_heading(doc, "时间对齐模型", 2)
    add_body(doc, "设第 k 个文本词对应时间区间 Iₖ 等于从 sₖ 到 eₖ。对语音或视觉原始帧特征 φₘ(t)，使用帧质量权重 qₘ(t) 在区间内进行加权池化。")
    add_omml_equation(doc, "xₖ⁽ᵐ⁾ = [Σₜ∈Iₖ qₘ(t)φₘ(t)] / [Σₜ∈Iₖ qₘ(t) + ε]", "式 2")
    add_body(doc, "完成对齐后，每条样本形成三组共享序列长度的特征。")
    add_omml_equation(doc, "X⁽ᵀ⁾ ∈ ℝᴸˣ⁷⁶⁸    X⁽ᴬ⁾ ∈ ℝᴸˣ⁷⁴    X⁽ⱽ⁾ ∈ ℝᴸˣ³⁵", "式 3")
    add_heading(doc, "长短序列处理", 2)
    add_bullets(doc, [
        "短于 50 个位置的样本在右侧填充，并保存原始有效长度。",
        "超过 50 个位置时，优先将相邻词聚合为 50 个连续语义区间，避免直接删除视频后半段。",
        "如采用截断，必须保存截断标志、截断比例和被删除的时间范围。",
        "强制对齐失败时可以按词长和音频时长进行比例分配，但必须记录为近似对齐。",
    ])
    add_heading(doc, "特征文件规范", 2)
    add_caption(doc, "表 9 问题一每条样本应保存的字段")
    add_table(doc, ["字段组", "字段", "用途"], [
        ["身份信息", "id、video_id、clip_id、raw_text", "确保视频、标签和特征一一对应"],
        ["时间信息", "duration、interval_start、interval_end", "把序列位置映射回原始素材"],
        ["特征", "text、audio、vision", "提供统一数值输入"],
        ["掩码", "padding_mask、availability_mask、quality_mask", "区分填充、缺失和质量失效"],
        ["质量记录", "face_failure_rate、alignment_confidence、truncated", "支持异常分析和复现"],
        ["环境记录", "工具版本、模型名称、参数和随机种子", "满足可复现性要求"],
    ], widths=[2.4, 5.2, 6.0], font_size=9.2)
    add_body(doc, "典型样本验证应在同一时间轴上展示英文词、语音波形或能量、视频关键帧、三个模态的特征位置以及最终标签。该图用于证明对齐规则可以核查，而不是只展示抽象网络结构。")

    add_heading(doc, "五 问题二缺失条件下鲁棒预测", 1, page_break=True)
    add_heading(doc, "推荐模型结构", 2)
    add_body(doc, "建议使用规模适中的缺失感知多模态 Transformer。4850 条样本不足以从头稳定训练大型扩散模型。每个模态先通过独立投影和两层时序编码器映射到统一隐藏维度，再由跨模态模块重建被遮挡位置，最后使用动态可靠性权重完成融合。隐藏维度可以从 128 起步，注意力头数取 4，dropout 在 0.1 至 0.3 之间由验证集选择。")
    add_body(doc, "设 Hₘ 为模态 m 的时序编码。缺失位置的表示由其他模态和本模态邻域共同预测，得到 H 波浪。融合权重由内容、掩码和重建不确定性共同决定。")
    add_omml_equation(doc, "αₘ,ₜ = exp(rₘ,ₜ)Mₘ,ₜ / [Σⱼ exp(rⱼ,ₜ)Mⱼ,ₜ + ε]", "式 4")
    add_omml_equation(doc, "Hₜ = Σₘ αₘ,ₜ H̃ₘ,ₜ", "式 5")
    add_body(doc, "融合序列经过注意力池化后进入两个独立输出头。分类头输出 Negative、Neutral 和 Positive 三类概率，回归头输出 -3 至 3 的连续情感强度。独立分类头十分必要，因为连续回归模型几乎不会精确输出 0，若直接按回归值正负划分类别，中性类性能会显著下降。")

    add_heading(doc, "连续区间缺失增强", 2)
    add_body(doc, "训练阶段应生成与附件三相似的连续缺失块，不能只随机删除离散时间点。人工遮挡只能施加在原本有效的位置，重建损失也只计算在这些被人工遮挡的位置上。")
    add_caption(doc, "表 10 缺失实验因素")
    add_table(doc, ["因素", "建议水平", "分析目的"], [
        ["缺失模态", "T、A、V、TA、TV、AV、TAV", "比较单模态和多模态缺失"],
        ["缺失比例", "10%、20%、30%、40%、50%", "估计性能下降斜率"],
        ["缺失位置", "开头、中间、结尾、随机", "判断情感关键区域的位置效应"],
        ["随机重复", "每种条件 5 个随机种子", "估计均值、方差和置信区间"],
    ], widths=[2.6, 5.0, 6.0])
    add_heading(doc, "联合目标函数", 2)
    add_omml_equation(doc, "L = λcLCE + λrLSmoothL1 + λrecLreconstruct + λconLconsistency", "式 6")
    add_bullets(doc, [
        "LCE 使用类别加权交叉熵，缓解正向样本占比偏高的问题。",
        "LSmoothL1 用于情感强度回归，对少量极端标签比均方误差更稳定。",
        "Lreconstruct 只约束人工遮挡的原始有效位置。",
        "Lconsistency 约束完整输入与缺失输入的预测保持接近。",
    ])
    add_heading(doc, "缺失影响规律", 2)
    add_body(doc, "附件三没有标签，不能用来分析性能下降。缺失规律必须在附件二 valid 和 test 的有标签样本上通过配对模拟得到。对同一条样本比较完整输入和缺失输入，可以消除样本难度差异。")
    add_omml_equation(doc, "ΔM = β₀ + βmodality + β₁ρ + βposition + βinteraction + ε", "式 7")
    add_body(doc, "其中 ρ 为缺失比例，ΔM 表示 Accuracy、F1、MAE 或 Pearson 的变化。可以使用重复测量方差分析或混合效应回归检验模态类型、缺失比例、缺失位置及其交互项。正文应同时给出均值、95% bootstrap 置信区间和缺失比例曲线。")

    add_heading(doc, "六 问题三可解释性预测", 1, page_break=True)
    add_heading(doc, "模态贡献", 2)
    add_body(doc, "问题三可以复用问题二训练得到的缺失感知模型。该模型在训练中见过不同模态子集，因此用它评估模态移除后的输出变化比临时向普通模型输入全零模态更可靠。三个模态只有 8 种组合，可以计算精确 Shapley 值，无需采样近似。")
    add_omml_equation(doc, "φₘ = ΣS⊆M∖{m} [|S|!(|M|−|S|−1)!/|M|!] · [f(S∪{m})−f(S)]", "式 8")
    add_body(doc, "对分类任务使用预测类别的 logit 计算贡献，对回归任务使用预测强度计算贡献。若某个模态的 Shapley 值为负，说明该模态与最终判断存在冲突，应保留贡献符号，不能强行归一化为正数。")
    add_heading(doc, "局部证据定位", 2)
    add_body(doc, "在每个模态上滑动连续窗口，依次遮挡长度为 2 至 5 的位置，并重新计算预测。某窗口被遮挡后引起的输出变化越大，该窗口越重要。")
    add_omml_equation(doc, "Iₘ,[i,j] = |f(X) − f(X without modality m on interval [i,j])|", "式 9")
    add_bullets(doc, [
        "文本位置通过 tokenizer offset 映射回英文短语。",
        "语音位置输出起止时间，建议精确到 0.1 秒。",
        "视觉位置输出起止时间和中点关键帧。",
        "对齐特征没有直接保存真实时间戳时，应对附件四视频和转写重新执行强制对齐。",
    ])
    add_heading(doc, "解释可信度", 2)
    add_body(doc, "注意力权重只能作为辅助可视化，不能单独作为解释。解释结果还应通过删除和保留实验验证。删除模型认为最关键的片段后，预测置信度或强度应明显下降；只保留关键片段时，预测应尽量保持。关键片段的影响还应显著大于同长度随机片段。")
    add_caption(doc, "表 11 附件四解释卡字段")
    add_table(doc, ["字段组", "建议字段", "说明"], [
        ["预测", "id、pred_class、pred_score、三类概率", "给出极性、强度和置信度"],
        ["模态作用", "shap_text、shap_audio、shap_vision、main_modality", "给出带符号贡献和主要参考模态"],
        ["文本证据", "key_text、token_start、token_end", "映射回原始英文短语"],
        ["语音证据", "audio_start、audio_end", "定位到可回听时间段"],
        ["视觉证据", "vision_start、vision_end、key_frame", "定位到视频帧和时间段"],
        ["可信度", "deletion_drop、sufficiency_score、quality_flag", "验证解释并记录模态质量"],
    ], widths=[2.3, 6.2, 5.3], font_size=9.1)

    add_heading(doc, "七 实验设计与评价方法", 1, page_break=True)
    add_heading(doc, "基础模型与消融", 2)
    add_caption(doc, "表 12 推荐实验矩阵")
    add_table(doc, ["实验组", "模型或改动", "目的"], [
        ["单模态基线", "Text only、Audio only、Vision only", "识别各模态单独贡献和数据质量"],
        ["融合基线", "早期拼接、固定权重晚期融合", "建立简单可复现参照"],
        ["稳健基线", "模态 dropout，不使用重建", "检验随机失活的作用"],
        ["完整模型", "连续缺失增强、重建、动态门控、多任务", "形成主要结果"],
        ["消融一", "去除连续缺失增强", "检验缺失分布匹配的作用"],
        ["消融二", "去除重建模块", "检验跨模态补偿"],
        ["消融三", "动态门控改为固定权重", "检验样本级可靠性估计"],
        ["消融四", "去除一致性或辅助任务", "检验联合目标的作用"],
    ], widths=[2.5, 6.0, 5.3], font_size=9.1)
    add_heading(doc, "指标计算", 2)
    add_bullets(doc, [
        "分类报告 Accuracy、Macro F1、Weighted F1 和混淆矩阵。",
        "回归报告 MAE、Pearson r，并绘制预测值与真实值散点图。",
        "缺失鲁棒性报告完整性能、缺失性能、绝对下降量和相对下降率。",
        "解释报告删除曲线、AOPC、充分性、必要性和随机片段对照。",
        "所有主要结果使用样本 bootstrap 给出 95% 置信区间。",
    ])
    add_heading(doc, "阈值和模型选择", 2)
    add_body(doc, "分类任务应直接使用三分类头。若需要将回归输出转为类别作辅助对照，应在 valid 上选择负向阈值 τ− 和正向阈值 τ+，区间内判为中性。任何阈值都不能在 test、附件三或附件四上调整。模型选择可优先考虑归一化综合指标，例如 Macro F1、Pearson r 和负 MAE 的加权和，但论文必须同时列出每个原始指标。")

    add_heading(doc, "八 结果文件与论文组织", 1, page_break=True)
    add_heading(doc, "专项结果文件", 2)
    add_caption(doc, "表 13 建议提交字段")
    add_table(doc, ["文件", "必须字段", "建议附加字段"], [
        ["附件三预测 CSV", "id、pred_class、pred_score", "p_negative、p_neutral、p_positive、missing_type、missing_ratio"],
        ["附件四解释 CSV", "id、pred_class、pred_score、main_modality、模态作用程度、关键证据", "三类概率、关键时间、关键帧、解释可信度、质量标志"],
    ], widths=[3.0, 5.8, 5.8], font_size=9.1)
    add_body(doc, "附件三文件中没有显式 id 时，统一使用文件名去掉扩展名作为样本编号。所有数值字段应明确精度和范围，分类名称统一使用 Negative、Neutral 和 Positive，避免同一文件中混用中文、英文和数字编码。")
    add_heading(doc, "论文正文结构", 2)
    add_numbered(doc, [
        "摘要，包括三个问题的方法、核心结果和主要结论。",
        "问题重述与总体技术路线。",
        "数据说明、探索性分析和质量异常。",
        "问题一的特征定义、时间对齐模型、典型样本和复现规则。",
        "问题二的网络结构、损失函数、训练方案、缺失规律和附件三结果。",
        "问题三的模态贡献、局部证据、解释可信度和附件四结果。",
        "模型评价、消融实验、误差归因和局限性。",
        "结论与可复现性说明。",
    ])
    add_heading(doc, "附件大小控制", 2)
    add_body(doc, "竞赛附件总大小不超过 50 MB。问题一特征建议使用 float32 和压缩存储；融合模型隐藏维度控制在 128 左右；冻结 BERT 的完整预训练参数不宜重复提交，可在说明文档中记录模型名称、版本和校验信息，只提交本题训练得到的投影层、融合层和预测头参数。提交前应实际核对竞赛对外部预训练权重获取方式的具体要求。")

    add_heading(doc, "九 实施顺序与风险控制", 1, page_break=True)
    add_caption(doc, "表 14 推荐实施顺序")
    add_table(doc, ["阶段", "主要工作", "完成标志"], [
        ["第一阶段", "数据读取、ID 核对、标签统计、三类掩码生成", "所有样本可读取，异常有清单"],
        ["第二阶段", "问题一特征提取和典型样本对齐图", "100 条特征与时间记录齐全"],
        ["第三阶段", "单模态和简单融合基线", "四项官方指标可重复"],
        ["第四阶段", "连续缺失增强、重建和动态门控", "缺失曲线优于融合基线"],
        ["第五阶段", "Shapley 值、窗口遮挡和解释验证", "20 条解释卡可回看"],
        ["第六阶段", "消融、误差分析、结果文件和论文整理", "表图与提交文件一致"],
    ], widths=[2.4, 6.8, 4.6], font_size=9.2)
    add_heading(doc, "主要风险", 2)
    add_bullets(doc, [
        "把全零视觉特征当作真实中性表情。处理方式是独立设置质量掩码。",
        "训练使用 text，附件三推理使用 text_bert 临时生成表示。处理方式是统一从 text_bert 编码。",
        "把附件三用于调参。附件三无标签，只能做最终推理。",
        "使用回归值正负号直接生成三分类。处理方式是设置独立分类头。",
        "只展示注意力热力图，不验证解释。处理方式是加入删除、保留和随机片段对照。",
        "未对齐版盲目信任 vision_lengths。处理方式是联合声明长度、非零位置和质量规则。",
        "模型规模过大导致小样本过拟合或附件超限。处理方式是冻结文本骨干并控制融合层规模。",
    ])
    add_heading(doc, "最终建议", 2)
    add_body(doc, "最稳妥的竞赛方案是以 aligned_50 为主，先完成数据掩码和轻量基线，再加入连续缺失增强、跨模态重建和动态可靠性融合。问题三直接复用该缺失感知模型，以精确模态 Shapley 值解释模态差异，以连续窗口遮挡定位证据，并用删除实验验证解释。整篇论文应突出数据质量、缺失机制与解释可信度之间的逻辑关系，而不是仅强调网络结构复杂度。")

    add_heading(doc, "数据来源说明", 1, page_break=True)
    add_body(doc, "本文档中的题目要求、字段定义和数据形状来自赛题文档。样本数量、类别分布和数据异常来自对附件一标签表、附件二标签表以及 aligned_50.pkl、unaligned_50.pkl、附件三和附件四特征文件的只读核查。建模阶段不得引入其他情感数据集参与训练、微调、阈值选择或结果统计。")

    # Ensure all ordinary runs use intended fonts.
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run._element.getparent().tag != qn("m:oMath"):
                if paragraph.style.name == "Title":
                    set_run_font(run, HEAD_CN, 22, bold=True)
                elif paragraph.style.name.startswith("Heading"):
                    level = int(paragraph.style.name.split()[-1])
                    set_run_font(run, HEAD_CN, {1: 16, 2: 13.5, 3: 12}.get(level, 11), bold=True)
                elif paragraph.style.name == "Subtitle":
                    set_run_font(run, BODY_CN, 12)
                elif run.font.size is None:
                    set_run_font(run)

    doc.core_properties.title = "复杂场景下多模态情感预测建模详细解题方案"
    doc.core_properties.subject = "数学建模竞赛 E 题解题方案"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "多模态情感预测 数学建模 模态缺失 可解释性"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
